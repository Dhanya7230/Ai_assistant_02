"""
Extract plain text from uploaded legal documents.

Supports .txt, .pdf, .docx. Pure functions over bytes -> str, so they're
easy to unit test without spinning up the web server.
"""
from __future__ import annotations

import io

from backend.security import validate_extracted_text


class ParsingError(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        text = _extract_txt(content)
    elif ext == "pdf":
        text = _extract_pdf(content)
    elif ext == "docx":
        text = _extract_docx(content)
    else:
        raise ParsingError(f"Unsupported extension: .{ext}")
    return validate_extracted_text(text)


def _extract_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="replace")


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParsingError("pypdf is not installed. Run: pip install pypdf") from exc

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ParsingError(f"Could not open PDF: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise ParsingError("This PDF is password-protected.")

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ParsingError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ParsingError(f"Could not open DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def chunk_text(text: str, max_chars: int = 12_000, overlap: int = 500) -> list[str]:
    """Split long text into overlapping chunks for map-reduce style processing."""
    if len(text) <= max_chars:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks
